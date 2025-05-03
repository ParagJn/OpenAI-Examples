"""
document_utils.py

Utility functions for handling document reading and image extraction from docx files.
Provides reusable helpers for extracting text and images from uploaded files.
"""

import os

def ensure_image_folder():
    """
    Ensures the image folder exists for saving extracted images.

    Returns:
        str: The absolute path to the images folder.
    """
    img_folder = os.path.join(os.getcwd(), "document", "images")
    os.makedirs(img_folder, exist_ok=True)
    return img_folder

def extract_images_from_docx(file, img_folder):
    """
    Extracts images from a docx file and saves them to the specified folder.

    Args:
        file: The uploaded docx file object.
        img_folder (str): Path to the folder where images will be saved.

    Returns:
        list: List of file paths to the saved images.
    """
    from docx import Document
    images = []
    doc = Document(file)
    rels = doc.part.rels
    for rel in rels:
        rel_obj = rels[rel]
        if "image" in rel_obj.target_ref:
            img_part = rel_obj.target_part
            img_data = img_part.blob
            img_name = os.path.basename(img_part.partname)
            img_path = os.path.join(img_folder, img_name)
            with open(img_path, "wb") as f:
                f.write(img_data)
            images.append(img_path)
    return images

def read_file(file):
    """
    Reads the uploaded file and extracts text and images (if docx).

    Args:
        file: The uploaded file object.

    Returns:
        tuple: (document_text, images_list)
    """
    if file.type == "text/plain":
        return file.read().decode("utf-8"), []
    elif file.type in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document"]:
        import docx
        img_folder = ensure_image_folder()
        images = extract_images_from_docx(file, img_folder)
        doc = docx.Document(file)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text, images
    else:
        return None, []
